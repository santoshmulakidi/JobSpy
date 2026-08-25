from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "local-copilot" / "Local_Windows_AI_Copilot_Product_Dossier.docx"
SOURCES = [
    ROOT / "docs" / "local-copilot" / "00-research-dossier.md",
    ROOT / "docs" / "local-copilot" / "01-business-requirements-document.md",
    ROOT / "docs" / "local-copilot" / "02-system-architecture.md",
    ROOT / "docs" / "local-copilot" / "03-security-privacy-threat-model.md",
    ROOT / "docs" / "local-copilot" / "04-test-and-validation-strategy.md",
    ROOT / "docs" / "superpowers" / "plans" / "2026-08-24-local-windows-ai-copilot.md",
]

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_GRAY = "F2F4F7"
MUTED = "666666"
INK = "1B1F23"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_font(run, name="Calibri", size=11, color=INK, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_field(paragraph) -> None:
    paragraph.add_run("Page ")
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    code = doc.styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Inches(0.15)
    code.paragraph_format.right_indent = Inches(0.15)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.0
    p_pr = code._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F6F8FA")
    p_pr.append(shd)


def add_inline(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Consolas", size=9.5, color=DARK_BLUE)
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(f"{label} ({url})")
            run.font.color.rgb = RGBColor.from_string(BLUE)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    widths = [9360 // cols] * cols
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    mark_table_header(table.rows[0])
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.text = row[c_idx].strip() if c_idx < len(row) else ""
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    set_font(run, size=9.3, bold=(r_idx == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_markdown(doc: Document, path: Path, skip_first_h1: bool = False) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    idx = 0
    in_code = False
    code_lines: list[str] = []
    para_lines: list[str] = []
    first_h1_seen = False

    def flush_para() -> None:
        nonlocal para_lines
        if para_lines:
            p = doc.add_paragraph()
            add_inline(p, " ".join(line.strip() for line in para_lines))
            para_lines = []

    while idx < len(lines):
        line = lines[idx]
        if line.startswith("```"):
            flush_para()
            if in_code:
                p = doc.add_paragraph(style="Code Block")
                p.add_run("\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            idx += 1
            continue
        if in_code:
            code_lines.append(line)
            idx += 1
            continue
        if line.startswith("|") and idx + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[idx + 1]):
            flush_para()
            table_rows = [[cell.strip() for cell in line.strip("|").split("|")]]
            idx += 2
            while idx < len(lines) and lines[idx].startswith("|"):
                table_rows.append([cell.strip() for cell in lines[idx].strip("|").split("|")])
                idx += 1
            add_markdown_table(doc, table_rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            flush_para()
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and not first_h1_seen:
                first_h1_seen = True
                if skip_first_h1:
                    idx += 1
                    continue
            doc.add_heading(text, level=level)
            idx += 1
            continue
        bullet = re.match(r"^\s*-\s+(.*)$", line)
        number = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if bullet or number:
            flush_para()
            p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline(p, (bullet or number).group(1).strip())
            idx += 1
            continue
        if not line.strip():
            flush_para()
            idx += 1
            continue
        if line.strip() == "---":
            flush_para()
            idx += 1
            continue
        para_lines.append(line)
        idx += 1
    flush_para()


def build() -> None:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("LOCAL WINDOWS AI COPILOT  |  PRODUCT DOSSIER")
    set_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(footer)
    for run in footer.runs:
        set_font(run, size=8.5, color=MUTED)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(28)
    kicker = doc.add_paragraph()
    run = kicker.add_run("PRODUCT RESEARCH  |  BUSINESS REQUIREMENTS  |  TECHNICAL DESIGN")
    set_font(run, size=9.5, color=BLUE, bold=True)
    kicker.paragraph_format.space_after = Pt(8)

    title = doc.add_paragraph()
    run = title.add_run("Local Windows AI Copilot")
    set_font(run, size=26, color=INK, bold=True)
    title.paragraph_format.space_after = Pt(5)

    subtitle = doc.add_paragraph()
    run = subtitle.add_run("Clean-room product dossier for a secure, local-first Electron implementation")
    set_font(run, size=14, color=MUTED)
    subtitle.paragraph_format.space_after = Pt(22)

    metadata = [
        ("Status", "Approved design baseline"),
        ("Platform", "Windows 10/11 x64"),
        ("Architecture", "Electron + React + TypeScript"),
        ("Data model", "Local-first, bring your own API key"),
        ("Prepared", "24 August 2026"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"{label}: ")
        set_font(r1, size=10.5, bold=True)
        r2 = p.add_run(value)
        set_font(r2, size=10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    note = doc.add_table(rows=1, cols=1)
    note.style = "Table Grid"
    set_table_geometry(note, [9360])
    mark_table_header(note.rows[0])
    set_cell_shading(note.cell(0, 0), "E8EEF5")
    p = note.cell(0, 0).paragraphs[0]
    r = p.add_run("Scope note. This document describes an independent product using public APIs and clean-room observations. It does not include copied credentials, branding, private services, or proprietary assets.")
    set_font(r, size=10, color=DARK_BLUE, bold=True)

    doc.add_page_break()
    doc.add_heading("Document map", level=1)
    for label, description in (
        ("Part I - Research dossier", "Observed CueFlow architecture, models, storage, and security lessons."),
        ("Part II - Business requirements", "MVP scope, users, requirements, success metrics, and release acceptance."),
        ("Part III - System architecture", "Processes, interfaces, data, networking, overlay, and error handling."),
        ("Part IV - Security and privacy", "Trust boundaries, threats, controls, retention, and release gates."),
        ("Part V - Test strategy", "Automated, compatibility, performance, privacy, and release validation."),
        ("Part VI - Implementation plan", "Twelve test-driven tasks leading to a signed Windows MVP."),
    ):
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(f"{label}: ")
        r.bold = True
        p.add_run(description)

    part_titles = [
        "Part I - Research dossier",
        "Part II - Business requirements",
        "Part III - System architecture",
        "Part IV - Security and privacy threat model",
        "Part V - Test and validation strategy",
        "Part VI - Implementation plan",
    ]
    for part_title, source in zip(part_titles, SOURCES):
        doc.add_page_break()
        doc.add_heading(part_title, level=1)
        parse_markdown(doc, source, skip_first_h1=True)

    doc.core_properties.title = "Local Windows AI Copilot Product Dossier"
    doc.core_properties.subject = "Research, BRD, architecture, security, testing, and implementation plan"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = "Electron, Windows, AI copilot, BRD, architecture, privacy"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
